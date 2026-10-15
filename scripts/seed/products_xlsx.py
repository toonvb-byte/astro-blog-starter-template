"""Product collection sheet <-> finder-data.json.

  python3 scripts/seed/products_xlsx.py export out.xlsx   # data.json -> sheet (prefilled with current products)
  python3 scripts/seed/products_xlsx.py import in.xlsx    # sheet -> data.json (rows with Status = Ready)

After import: python3 scripts/seed/build_finder_seed.py > scripts/seed/goal_finder_seed.sql
"""
import json, os, re, sys
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

HERE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(HERE, "finder-data.json")
SLOTS = 12         # ingredient slots per product
EMPTY_ROWS = 150   # blank rows prepared for new products

BASE = [  # (key, header, width, help)
    ("status", "Status", 10, "Draft = still collecting. Ready = import into the database."),
    ("id", "Product ID", 24, "Short unique code, lowercase-with-dashes, e.g. solgar-zinc-22mg. Leave empty for new products: it is created from brand + name."),
    ("brand", "Brand", 18, "As on the pack, e.g. Solgar"),
    ("name", "Product name", 32, "Name as sold, incl. strength, e.g. Zinc Picolinate 22 mg"),
    ("ean", "EAN barcode", 16, "13-digit barcode (on pack or pharmacy page). Used to match shop feeds automatically. Type as text."),
    ("ES", "Spain", 10, "verified = you found it on a Spanish shop. likely = brand sold in Spain, no listing checked. Empty = not sold."),
    ("DE", "Germany", 10, "Same as Spain, for German shops."),
    ("pack", "Pack size", 14, "e.g. 90 caps, 250 g, 200 ml"),
    ("servings", "Servings per pack", 10, "Number of daily doses in the pack (whole number)"),
    ("dose_en", "Daily dose (EN)", 26, "e.g. 2 capsules a day"),
    ("dose_es", "Daily dose (ES)", 26, "e.g. 2 cápsulas al día"),
    ("vegan", "Vegan", 8, "Yes / No / empty if unknown"),
    ("source", "Label source URL", 34, "Page where you read the ingredient amounts"),
    ("price", "Price €", 9, "Indicative price incl. VAT"),
    ("price_where", "Price at shop", 16, "e.g. farma2go.com"),
    ("price_date", "Price checked", 12, "Date you saw the price"),
    ("partial", "Label incomplete?", 10, "Yes if the label hides amounts or you could not capture everything"),
    ("flag_en", "Warning (EN)", 30, "Optional warning shown on the product card"),
    ("flag_es", "Warning (ES)", 30, "Same warning in Spanish"),
    ("excluded", "Exclude because", 24, "Fill in to keep it OUT of recommendations, e.g. registered medicine"),
    ("notes", "Your notes", 24, "Not imported. Anything for yourself."),
]
SLOT_COLS = [("ing", "Ingredient {n}", 24), ("amt", "Amount {n}", 10), ("unit", "Unit {n}", 8), ("form", "Form {n}", 18)]

ARIAL = "Arial"
F_HEAD = Font(name=ARIAL, bold=True, color="FFFFFF", size=10)
F_BODY = Font(name=ARIAL, size=10)
F_AUTO = Font(name=ARIAL, size=10, color="5A6863", italic=True)
FILL_HEAD = PatternFill("solid", fgColor="0B6B4F")
FILL_SLOT = PatternFill("solid", fgColor="0B4F3B")
FILL_AUTO = PatternFill("solid", fgColor="EEF1EE")
FILL_INPUT = PatternFill("solid", fgColor="FFFBE6")
THIN = Side(style="thin", color="DCE1DD")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def columns():
    cols = [(k, h, w, help_) for k, h, w, help_ in BASE]
    for n in range(1, SLOTS + 1):
        for k, h, w in SLOT_COLS:
            cols.append((f"{k}{n}", h.format(n=n), w, None))
    return cols


def export(path):
    d = json.load(open(DATA))
    ings = d["ingredients"]
    wb = Workbook()

    # --- How to fill in ---------------------------------------------------
    hw = wb.active
    hw.title = "How to fill in"
    lines = [
        ("NutriMetrics product collection sheet", True),
        ("", False),
        ("1. Add one product per row on the 'Products' tab. The first 30 rows are the products already in the database: use them as examples, and add missing EANs or prices.", False),
        ("2. Yellow cells are for you to fill in. Grey 'Unit' cells fill themselves when you pick an ingredient.", False),
        ("3. Ingredients: pick from the dropdown (list on the 'Ingredients' tab). Enter the amount PER DAILY DOSE, in the unit shown, not per capsule.", False),
        ("   Example: 2 capsules a day of 150 mg L-theanine -> Amount = 300.", False),
        ("   If the label does not say the amount, leave Amount empty and set 'Label incomplete?' to Yes.", False),
        ("4. Only ingredients that matter for the goals need to be entered. Fillers, capsule shell etc. can be skipped.", False),
        ("5. Missing ingredient? Add it at the bottom of the 'Ingredients' tab (ID, names, unit). It is imported but only scores once its evidence is reviewed.", False),
        ("6. Set Status to 'Ready' when a row is complete. Only Ready rows are imported; Draft rows are ignored.", False),
        ("7. When done, send the file to Claude (or Google Sheets: File > Download > .xlsx) and it is imported into the database.", False),
        ("", False),
        ("Hover over a column header for an explanation of that column.", False),
    ]
    for i, (t, bold) in enumerate(lines, 1):
        c = hw.cell(row=i, column=1, value=t)
        c.font = Font(name=ARIAL, size=14 if i == 1 else 10, bold=bold)
    hw.column_dimensions["A"].width = 130

    # --- Ingredients reference --------------------------------------------
    iw = wb.create_sheet("Ingredients")
    ih = ["ID", "Name (EN)", "Name (ES)", "Unit", "Effective daily amount", "Safe upper limit / day"]
    for j, h in enumerate(ih, 1):
        c = iw.cell(row=1, column=j, value=h); c.font = F_HEAD; c.fill = FILL_HEAD
    for r, i in enumerate(ings, 2):
        vals = [i["id"], i["en"], i["es"], i["unit"], i.get("effective"), i.get("upper")]
        for j, v in enumerate(vals, 1):
            c = iw.cell(row=r, column=j, value=v); c.font = F_BODY; c.border = BORDER
    for j, w in enumerate([18, 28, 30, 9, 20, 20], 1):
        iw.column_dimensions[get_column_letter(j)].width = w
    iw.freeze_panes = "A2"
    note_row = len(ings) + 3
    iw.cell(row=note_row, column=1, value="Add new ingredients below this list (ID, names, unit). Effective amount is filled in after evidence review.").font = F_AUTO
    # name list for dropdown: B2:B200 (room for new rows)

    # --- Products ---------------------------------------------------------
    pw = wb.create_sheet("Products", 1)
    cols = columns()
    for j, (k, h, w, help_) in enumerate(cols, 1):
        c = pw.cell(row=1, column=j, value=h)
        c.font = F_HEAD
        c.fill = FILL_SLOT if re.match(r"(ing|amt|unit|form)\d", k) else FILL_HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = BORDER
        if help_:
            c.comment = Comment(help_, "NutriMetrics")
        pw.column_dimensions[get_column_letter(j)].width = w
    pw.row_dimensions[1].height = 32
    pw.freeze_panes = "E2"
    col = {k: j for j, (k, *_ ) in enumerate(cols, 1)}

    name_by_id = {i["id"]: i["en"] for i in ings}
    flagged = {}
    rows = []
    for p in d["products"]:
        r = {
            "status": "Ready", "id": p["id"], "brand": p["brand"], "name": p["name"], "ean": p.get("ean"),
            "ES": (p.get("markets") or {}).get("ES"), "DE": (p.get("markets") or {}).get("DE"),
            "pack": p.get("pack"), "servings": p.get("servings"),
            "dose_en": (p.get("dose") or {}).get("en"), "dose_es": (p.get("dose") or {}).get("es"),
            "vegan": "Yes" if p.get("vegan") else None, "source": p.get("source"),
            "price": (p.get("price") or {}).get("amt"), "price_where": (p.get("price") or {}).get("where"),
            "price_date": d["meta"]["checked"] if p.get("price") else None,
            "partial": "Yes" if p.get("partial") else None,
            "flag_en": " | ".join(f["en"] for f in p.get("flags", [])) or None,
            "flag_es": " | ".join(f["es"] for f in p.get("flags", [])) or None,
            "excluded": p.get("excluded"),
        }
        for n, x in enumerate(p["ingredients"][:SLOTS], 1):
            r[f"ing{n}"] = name_by_id[x["i"]]; r[f"amt{n}"] = x.get("amt"); r[f"form{n}"] = x.get("form")
        rows.append(r)

    last = 1 + len(rows) + EMPTY_ROWS
    for ri in range(2, last + 1):
        r = rows[ri - 2] if ri - 2 < len(rows) else {}
        for k, h, w, help_ in cols:
            c = pw.cell(row=ri, column=col[k])
            c.font = F_BODY; c.border = BORDER
            if k.startswith("unit"):
                n = k[4:]
                ing_cell = f"{get_column_letter(col['ing' + n])}{ri}"
                c.value = f'=IF({ing_cell}="","",IFERROR(INDEX(Ingredients!$D$2:$D$200,MATCH({ing_cell},Ingredients!$B$2:$B$200,0)),"?"))'
                c.fill = FILL_AUTO; c.font = F_AUTO
                continue
            v = r.get(k)
            if k == "ean" and v is not None:
                v = str(v)
            c.value = v
            c.fill = FILL_INPUT
            if k == "ean":
                c.number_format = "@"
            if k == "price":
                c.number_format = '#,##0.00 "€"'
        if not r:
            pw.cell(row=ri, column=col["status"], value="Draft")

    def dv(formula, colkey, prompt=None):
        v = DataValidation(type="list", formula1=formula, allow_blank=True, showDropDown=False)
        if prompt:
            v.error = prompt; v.errorTitle = "Invalid value"; v.showErrorMessage = True
        pw.add_data_validation(v)
        L = get_column_letter(col[colkey]); v.add(f"{L}2:{L}{last}")

    dv('"Draft,Ready"', "status", "Choose Draft or Ready")
    dv('"verified,likely"', "ES", "Choose verified, likely, or leave empty")
    dv('"verified,likely"', "DE", "Choose verified, likely, or leave empty")
    dv('"Yes,No"', "vegan")
    dv('"Yes,No"', "partial")
    for n in range(1, SLOTS + 1):
        dv("=Ingredients!$B$2:$B$200", f"ing{n}", "Pick an ingredient from the list (or add it to the Ingredients tab first)")
    for n in range(1, SLOTS + 1):
        v = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
        v.error = "Amount must be a number (per daily dose)"; v.showErrorMessage = True
        pw.add_data_validation(v); L = get_column_letter(col[f"amt{n}"]); v.add(f"{L}2:{L}{last}")
    pw.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{last}"

    wb.active = 1
    wb.save(path)
    print(f"exported {len(rows)} products + {EMPTY_ROWS} empty rows -> {path}")


def slug(s):
    s = s.lower()
    for a, b in zip("áéíóúüñäöß", "aeiouunaos"):
        s = s.replace(a, b)
    return re.sub(r"[^a-z0-9]+", "-", s).strip("-")[:60]


def import_(path):
    d = json.load(open(DATA))
    wb = load_workbook(path, data_only=True)
    iw = wb["Ingredients"]
    known = {i["id"] for i in d["ingredients"]}
    for row in iw.iter_rows(min_row=2, values_only=True):
        iid, en, es, unit = (row + (None,) * 6)[:4]
        if not iid or not en or str(iid).startswith("Add new"):
            continue
        if iid not in known:
            d["ingredients"].append({"id": str(iid).strip(), "en": en, "es": es or en, "unit": unit or "mg", "effective": None})
            known.add(iid); print(f"new ingredient (needs evidence review): {iid}")
    by_name = {i["en"].lower(): i["id"] for i in d["ingredients"]}

    pw = wb["Products"]
    hdr = [c.value for c in pw[1]]
    cols = columns()
    key_by_header = {h: k for k, h, *_ in cols}
    keys = [key_by_header.get(h) for h in hdr]
    existing = {p["id"]: p for p in d["products"]}
    errors, imported = [], []
    for rn, row in enumerate(pw.iter_rows(min_row=2, values_only=True), 2):
        r = {k: v for k, v in zip(keys, row) if k}
        if (r.get("status") or "").strip().lower() != "ready":
            continue
        if not r.get("brand") or not r.get("name"):
            errors.append(f"row {rn}: brand and product name are required"); continue
        pid = r.get("id") or slug(f"{r['brand']}-{r['name']}")
        p = existing.get(pid, {"id": pid})
        p.update({"brand": r["brand"], "name": r["name"], "pack": r.get("pack"), "source": r.get("source")})
        if r.get("ean"):
            ean = re.sub(r"\D", "", str(r["ean"]).split(".")[0])
            if len(ean) not in (8, 12, 13, 14):
                errors.append(f"row {rn}: EAN '{r['ean']}' should have 8, 12, 13 or 14 digits")
            p["ean"] = ean
        if r.get("servings"):
            p["servings"] = int(r["servings"])
        if r.get("dose_en") or r.get("dose_es"):
            p["dose"] = {"en": r.get("dose_en") or r.get("dose_es"), "es": r.get("dose_es") or r.get("dose_en")}
        p["markets"] = {m: r[m].strip().lower() for m in ("ES", "DE") if r.get(m)}
        if not p["markets"] and not r.get("excluded"):
            errors.append(f"row {rn} ({pid}): no market set (Spain/Germany)")
        p["vegan"] = (r.get("vegan") or "").lower() == "yes" or None
        p["partial"] = (r.get("partial") or "").lower() == "yes" or None
        if r.get("price"):
            p["price"] = {"amt": float(r["price"]), "where": r.get("price_where") or ""}
        else:
            p.pop("price", None)
        if r.get("flag_en"):
            p["flags"] = [{"en": r["flag_en"], "es": r.get("flag_es") or r["flag_en"]}]
        else:
            p.pop("flags", None)
        if r.get("excluded"):
            p["excluded"] = r["excluded"]
        else:
            p.pop("excluded", None)
        ings = []
        for n in range(1, SLOTS + 1):
            nm = r.get(f"ing{n}")
            if not nm:
                continue
            iid = by_name.get(str(nm).lower())
            if not iid:
                errors.append(f"row {rn} ({pid}): unknown ingredient '{nm}'"); continue
            amt = r.get(f"amt{n}")
            if amt is not None and not isinstance(amt, (int, float)):
                errors.append(f"row {rn} ({pid}): amount for {nm} is not a number"); amt = None
            x = {"i": iid, "amt": amt}
            if r.get(f"form{n}"):
                x["form"] = r[f"form{n}"]
            ings.append(x)
        if not ings and not r.get("excluded"):
            errors.append(f"row {rn} ({pid}): no ingredients")
        p["ingredients"] = ings
        for k in [k for k, v in list(p.items()) if v is None]:
            del p[k]
        existing[pid] = p
        imported.append(pid)
    if errors:
        print("NOT IMPORTED — fix these first:"); print("\n".join("  " + e for e in errors)); sys.exit(1)
    d["products"] = list(existing.values())
    json.dump(d, open(DATA, "w"), ensure_ascii=False, indent=1)
    print(f"imported {len(imported)} ready rows; database now has {len(d['products'])} products")


if __name__ == "__main__":
    cmd, path = sys.argv[1], sys.argv[2]
    export(path) if cmd == "export" else import_(path)
